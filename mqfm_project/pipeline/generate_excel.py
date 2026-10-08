"""
pipeline/generate_excel.py — Step 3: Output Generation
----------------------------------------------------------
Appends validated rows into the automated MIS Excel log. Each document
type gets its own sheet (created on first use) inside the same workbook,
with columns driven entirely by that type's schema.
"""
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .schemas import get_type


def _get_or_create_workbook(path: str):
    if os.path.exists(path):
        return load_workbook(path)
    wb = Workbook()
    wb.remove(wb.active)  # sheets are created per document type as needed
    return wb


def _get_or_create_sheet(wb, doc_type: str):
    schema = get_type(doc_type)
    sheet_name = schema["sheet_name"][:31]  # Excel sheet name limit

    if sheet_name in wb.sheetnames:
        return wb[sheet_name], schema

    ws = wb.create_sheet(title=sheet_name)
    headers = [schema["labels"].get(f, f) for f in schema["fields"]] + [
        "Source Document", "Validation Status", "Processed At",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(14, len(header) + 2)
    return ws, schema


def append_rows(xlsx_path: str, doc_type: str, rows: list) -> str:
    wb = _get_or_create_workbook(xlsx_path)
    ws, schema = _get_or_create_sheet(wb, doc_type)

    field_order = schema["fields"] + ["source_document_id", "validation_status", "processed_at"]

    for record in rows:
        next_row = ws.max_row + 1
        for col, field in enumerate(field_order, start=1):
            value = record.get(field, "")
            if isinstance(value, list):
                value = ", ".join(value)
            cell = ws.cell(row=next_row, column=col, value=value)
            cell.font = Font(name="Arial", size=10.5)
            if field == "validation_status":
                cell.fill = PatternFill("solid", fgColor="C6EFCE" if value == "PASSED" else "FFC7CE")
                cell.font = Font(name="Arial", size=10.5, bold=True, color="006100" if value == "PASSED" else "9C0006")

    wb.save(xlsx_path)
    return xlsx_path
